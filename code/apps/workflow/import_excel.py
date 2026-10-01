"""
Workflow App — Import Excel massif de recommandations en DRAFT (Story 6.5 — FR7)

Module frère de services.py (même niveau dans apps/workflow/).
NE PAS convertir services.py en package — casserait les imports existants.

Couverture FR7 : import opérationnel DRAFT uniquement.
    - import_tag reste null (règle « date non passée » du clean() s'applique).
    - Distinct de Story 6.8 (import historique, import_tag="IMPORTED").

Architecture découplée :
    parse_workbook()             → list[RowDraft]   (adaptateur de format Excel)
    validate_rows()              → ImportReport      (cœur pur, sans écriture)
    create_recommendations_bulk() → ImportBatch      (cœur atomique)
    build_import_template()      → Workbook          (générateur dynamique)

Le cœur (validate_rows + create_recommendations_bulk) ne connaît que list[RowDraft].
Un futur adaptateur PDF/IA produira la même structure et réutilisera ce cœur.
"""
from __future__ import annotations

import re
from dataclasses import dataclass, field
from datetime import date, datetime
from typing import Any

from django.core.exceptions import ValidationError
from django.db import IntegrityError, transaction
from django.db.models import Q
from django.utils import timezone
from django.utils.translation import gettext_lazy as _

from apps.audit.models import AuditLog
from apps.users.models import Department

from .models import Deliverable, ImportBatch, Recommendation, RecommendationSource
from .validators import validate_file_size, validate_magic_bytes


# ── Constantes ────────────────────────────────────────────────────────────────

MAX_ROWS = 50
DATA_SHEET_NAME = "Données"
INSTRUCTIONS_SHEET_NAME = "Instructions"

# Colonnes du template Excel (ordre = index 0-based dans la feuille)
HEADERS = [
    "Référence",            # 0 — COL_REFERENCE
    "Source",               # 1 — COL_SOURCE
    "Criticité",            # 2 — COL_PRIORITY
    "Description",          # 3 — COL_DESCRIPTION
    "Date cible",           # 4 — COL_DUE_DATE
    "Titre (mission)",      # 5 — COL_MISSION_LABEL
    "Date de mission",      # 6 — COL_MISSION_DATE
    "Direction contrôlée",  # 7 — COL_CTRL_DEPT
    "Direction concernée",  # 8 — COL_DEPT
    "Observations",         # 9 — COL_OBSERVATIONS
    "Dossiers en anomalies",  # 10 — COL_ANOMALIES
    "Livrable(s) attendu(s)",  # 11 — COL_DELIVERABLES
]

COL_REFERENCE     = 0
COL_SOURCE        = 1
COL_PRIORITY      = 2
COL_DESCRIPTION   = 3
COL_DUE_DATE      = 4
COL_MISSION_LABEL = 5
COL_MISSION_DATE  = 6
COL_CTRL_DEPT     = 7
COL_DEPT          = 8
COL_OBSERVATIONS  = 9
COL_ANOMALIES     = 10
COL_DELIVERABLES  = 11

_VALID_PRIORITIES = set(Recommendation.Priority.values)


# ── Structures de données ──


@dataclass
class RowDraft:
    """Ligne normalisée extraite de la feuille Excel — aucune dépendance modèle Django."""
    row_number: int
    reference: str
    source_raw: str
    priority: str
    description: str
    due_date_raw: Any        # date | str | None (str = illisible → erreur)
    mission_label: str
    mission_date_raw: Any    # date | str | None
    controlled_department_raw: str
    department_raw: str
    observations: str
    anomalous_dossiers: str
    deliverables_raw: str    # Brut avant split sur ";"


@dataclass
class ValidRow:
    """Ligne validée, prête pour Recommendation(**data)."""
    data: dict               # Champs modèle (sans deliverables)
    deliverables_data: list[str]


@dataclass
class ImportReport:
    """Résultat de validate_rows : compteurs, lignes valides, erreurs par ligne."""
    total: int
    valid_rows: list[ValidRow] = field(default_factory=list)
    errors_by_row: dict[int, list[str]] = field(default_factory=dict)

    @property
    def error_count(self) -> int:
        return len(self.errors_by_row)

    @property
    def valid_count(self) -> int:
        return len(self.valid_rows)

    def error_type_summary(self) -> str:
        """Résumé lisible des types d'erreurs pour le mini-récap dans le template."""
        counts: dict[str, int] = {}
        for errors in self.errors_by_row.values():
            for err in errors:
                lo = err.lower()
                if "obligatoire manquant" in lo:
                    key = "champ(s) obligatoire(s) manquant(s)"
                elif "en doublon" in lo:
                    key = "doublon(s) intra-fichier"
                elif "déjà présente en base" in lo:
                    key = "référence(s) existante(s) en base"
                elif "dans le passé" in lo:
                    key = "date(s) dans le passé"
                elif "illisible" in lo:
                    key = "date(s) illisible(s)"
                elif "source" in lo:
                    key = "source(s) inconnue(s)/inactive(s)"
                elif "direction" in lo:
                    key = "direction(s) inexistante(s)/inactive(s)"
                elif "criticité" in lo:
                    key = "criticité(s) invalide(s)"
                else:
                    key = "autre(s) erreur(s)"
                counts[key] = counts.get(key, 0) + 1
        return " ; ".join(f"{v} {k}" for k, v in counts.items())


# ── Génération du modèle Excel (dynamique) ────────────────────────────────────


def build_import_template():
    """
    Génère le modèle Excel d'import de recommandations.

    Dynamique : sources et directions sont chargées depuis la base au moment du clic.
    Jamais un fichier figé (listes déroulantes seraient périmées).

    Returns:
        openpyxl.Workbook prêt à être sérialisé en .xlsx.
    """
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = openpyxl.Workbook()

    # ── Onglet caché pour les listes (évite la limite 255 car. des formules inline) ──
    ws_lists = wb.active
    ws_lists.title = "_Listes"
    ws_lists.sheet_state = "hidden"

    sources = list(RecommendationSource.objects.filter(is_active=True).order_by("code"))
    departments = list(Department.objects.filter(is_active=True).order_by("code"))
    priorities = list(Recommendation.Priority.values)

    for i, src in enumerate(sources, start=1):
        ws_lists.cell(row=i, column=1, value=src.code)
    for i, prio in enumerate(priorities, start=1):
        ws_lists.cell(row=i, column=2, value=prio)
    for i, dept in enumerate(departments, start=1):
        ws_lists.cell(row=i, column=3, value=dept.code)

    # ── Onglet 1 : Données ────────────────────────────────────────────────────
    ws_data = wb.create_sheet(title=DATA_SHEET_NAME, index=0)

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="D4692A")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    req_mark_font = Font(bold=True, color="FFFFFF")

    col_widths = [22, 16, 14, 44, 16, 32, 16, 26, 26, 36, 36, 44]
    for col_idx, (header, width) in enumerate(zip(HEADERS, col_widths), start=1):
        cell = ws_data.cell(row=1, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        ws_data.column_dimensions[get_column_letter(col_idx)].width = width
    ws_data.row_dimensions[1].height = 22
    ws_data.freeze_panes = "A2"

    # Format Date Excel sur colonnes Date cible (E=5) et Date de mission (G=7)
    date_col_letters = [
        get_column_letter(COL_DUE_DATE + 1),
        get_column_letter(COL_MISSION_DATE + 1),
    ]
    for col_letter in date_col_letters:
        for row in range(2, MAX_ROWS + 2):
            ws_data.cell(row=row, column=ord(col_letter) - ord("A") + 1).number_format = "YYYY-MM-DD"

    # Listes déroulantes (depuis la feuille _Listes)
    n_src = max(len(sources), 1)
    n_dept = max(len(departments), 1)
    n_prio = len(priorities)
    data_rows = f"2:{MAX_ROWS + 1}"

    def _dv(col_idx: int, formula: str, error_msg: str, error_title: str) -> DataValidation:
        dv = DataValidation(
            type="list",
            formula1=formula,
            allow_blank=True,
            showErrorMessage=True,
            error=error_msg,
            errorTitle=error_title,
        )
        col = get_column_letter(col_idx)
        dv.sqref = f"{col}2:{col}{MAX_ROWS + 1}"
        return dv

    ws_data.add_data_validation(_dv(
        COL_SOURCE + 1,
        f"'_Listes'!$A$1:$A${n_src}",
        "Sélectionnez une source dans la liste (onglet Instructions).",
        "Source invalide",
    ))
    ws_data.add_data_validation(_dv(
        COL_PRIORITY + 1,
        f"'_Listes'!$B$1:$B${n_prio}",
        "Valeurs acceptées : CRITIQUE, HAUTE, MOYENNE, FAIBLE.",
        "Criticité invalide",
    ))
    ws_data.add_data_validation(_dv(
        COL_CTRL_DEPT + 1,
        f"'_Listes'!$C$1:$C${n_dept}",
        "Sélectionnez une direction dans la liste (onglet Instructions).",
        "Direction invalide",
    ))
    ws_data.add_data_validation(_dv(
        COL_DEPT + 1,
        f"'_Listes'!$C$1:$C${n_dept}",
        "Sélectionnez une direction dans la liste (onglet Instructions).",
        "Direction invalide",
    ))

    # ── Onglet 2 : Instructions ───────────────────────────────────────────────
    ws_inst = wb.create_sheet(title=INSTRUCTIONS_SHEET_NAME)
    ws_inst.column_dimensions["A"].width = 65
    ws_inst.column_dimensions["B"].width = 48

    title_font = Font(bold=True, size=13, color="D4692A")
    section_font = Font(bold=True, size=11)
    header2_font = Font(bold=True)

    row = 1
    ws_inst.cell(row=row, column=1, value="Guide d'utilisation — Modèle d'import de recommandations Sentinel").font = title_font
    row += 2

    # Colonnes obligatoires
    ws_inst.cell(row=row, column=1, value="COLONNES OBLIGATOIRES (*)").font = section_font
    row += 1
    required_headers = HEADERS[:5]
    for h in required_headers:
        ws_inst.cell(row=row, column=1, value=f"  • {h}")
        row += 1
    row += 1

    # Sources valides
    ws_inst.cell(row=row, column=1, value="SOURCES VALIDES").font = section_font
    row += 1
    ws_inst.cell(row=row, column=1, value="Code").font = header2_font
    ws_inst.cell(row=row, column=2, value="Libellé").font = header2_font
    row += 1
    for src in sources:
        ws_inst.cell(row=row, column=1, value=src.code)
        ws_inst.cell(row=row, column=2, value=src.label)
        row += 1
    row += 1

    # Criticités
    ws_inst.cell(row=row, column=1, value="CRITICITÉS AUTORISÉES").font = section_font
    row += 1
    prio_labels = {
        "CRITIQUE": "Risque critique — traitement immédiat",
        "HAUTE": "Risque élevé — traitement prioritaire",
        "MOYENNE": "Risque modéré",
        "FAIBLE": "Risque faible",
    }
    for prio in priorities:
        ws_inst.cell(row=row, column=1, value=prio)
        ws_inst.cell(row=row, column=2, value=prio_labels.get(prio, ""))
        row += 1
    row += 1

    # Directions valides
    ws_inst.cell(row=row, column=1, value="DIRECTIONS VALIDES").font = section_font
    row += 1
    ws_inst.cell(row=row, column=1, value="Code").font = header2_font
    ws_inst.cell(row=row, column=2, value="Nom").font = header2_font
    row += 1
    for dept in departments:
        ws_inst.cell(row=row, column=1, value=dept.code)
        ws_inst.cell(row=row, column=2, value=dept.name)
        row += 1
    row += 1

    # Règles de saisie
    ws_inst.cell(row=row, column=1, value="RÈGLES DE SAISIE").font = section_font
    row += 1
    rules = [
        "Format de date : ISO AAAA-MM-JJ (ex : 2026-09-30). Utilisez le type Date Excel OU ce format texte exact.",
        "La Date cible ne peut pas être dans le passé (date ≥ aujourd'hui).",
        "Livrables : plusieurs intitulés séparés par « ; » (ex : Procédure KYC ; Formation agents).",
        "Référence : format libre, maximum 50 caractères, unique dans le système et dans le fichier.",
        "Direction contrôlée : direction auditée (fait l'objet du contrôle).",
        "Direction concernée : direction responsable de la mise en œuvre de la recommandation.",
        f"Maximum {MAX_ROWS} recommandations par import. Au-delà, découpez en plusieurs fichiers.",
        "N'utilisez jamais de macros (.xlsm). Enregistrez toujours en .xlsx.",
    ]
    for rule in rules:
        ws_inst.cell(row=row, column=1, value=f"  • {rule}")
        row += 1

    return wb


# ── Helpers internes ──────────────────────────────────────────────────────────


def _normalize_str(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _parse_date_cell(value: Any) -> date | str | None:
    """
    Retourne un date (openpyxl date native), tente l'ISO AAAA-MM-JJ pour les textes.
    Tout autre format texte retourne la valeur brute (→ "date illisible" dans validate_rows).
    """
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if value is None:
        return None
    raw = str(value).strip()
    if not raw:
        return None
    try:
        return date.fromisoformat(raw)
    except (ValueError, TypeError):
        return raw


def get_existing_references(rows: list[RowDraft]) -> set[str]:
    """
    Références déjà présentes en base parmi celles du fichier (requête bornée).

    Ne charge que les références du fichier (≤ MAX_ROWS), insensible à la casse,
    au lieu de scanner toute la table Recommendation. Appelé par la vue de
    prévisualisation ET par create_recommendations_bulk (re-check concurrence AC8).
    """
    file_refs = [r.reference for r in rows if r.reference]
    if not file_refs:
        return set()
    q = Q()
    for ref in file_refs:
        q |= Q(reference__iexact=ref)
    return set(
        Recommendation.all_objects.filter(q).values_list("reference", flat=True)
    )


# ── Parser (adaptateur de format) ─────────────────────────────────────────────


def parse_workbook(file) -> list[RowDraft]:
    """
    Lit le fichier Excel et retourne la liste des RowDraft normalisés.

    Valide le format côté serveur (magic bytes, taille) AVANT tout parsing (AC2).
    Les ValidationError levées ici sont des erreurs de FORMAT — elles s'affichent
    via un message dédié, distinct du rapport de validation ligne-à-ligne.

    Args:
        file: Fichier Django (InMemoryUploadedFile / TemporaryUploadedFile).

    Returns:
        list[RowDraft]: Lignes normalisées (en-tête exclue, vides ignorées).

    Raises:
        ValidationError: Erreur de format (XLSM, >6 Mo, corrompu, onglet manquant, >MAX_ROWS).
    """
    import openpyxl

    # Validation format (AC2) — avant tout parsing
    validate_file_size(file)
    validate_magic_bytes(file, original_filename=getattr(file, "name", ""))

    file.seek(0)
    try:
        wb = openpyxl.load_workbook(file, read_only=True, data_only=True)
    except Exception as exc:
        # Capture l'exception openpyxl originale avant de la remplacer par le
        # message utilisateur. Permet de distinguer BadZipFile, KeyError, etc.
        try:
            import sentry_sdk
            sentry_sdk.capture_exception(exc)
        except Exception:
            pass
        raise ValidationError(
            _("Format de fichier non reconnu — déposez le modèle .xlsx Sentinel.")
        )

    if DATA_SHEET_NAME not in wb.sheetnames:
        raise ValidationError(
            _(
                f'L\'onglet « {DATA_SHEET_NAME} » est introuvable dans le fichier. '
                "Utilisez le modèle fourni par Sentinel."
            )
        )

    ws = wb[DATA_SHEET_NAME]
    rows: list[RowDraft] = []

    for row_idx, row_cells in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        # Ignorer les lignes entièrement vides
        if all(cell is None or str(cell).strip() == "" for cell in row_cells):
            continue

        # Plafond MAX_ROWS (AC7) — compter avant d'ajouter
        if len(rows) >= MAX_ROWS:
            raise ValidationError(
                _(
                    f"Le modèle accepte au maximum {MAX_ROWS} recommandations par import. "
                    "Votre fichier en contient davantage. Découpez l'import en plusieurs fichiers."
                )
            )

        def _col(idx: int) -> Any:
            try:
                return row_cells[idx]
            except IndexError:
                return None

        rows.append(RowDraft(
            row_number=row_idx,
            reference=_normalize_str(_col(COL_REFERENCE)),
            source_raw=_normalize_str(_col(COL_SOURCE)),
            priority=_normalize_str(_col(COL_PRIORITY)).upper(),
            description=_normalize_str(_col(COL_DESCRIPTION)),
            due_date_raw=_parse_date_cell(_col(COL_DUE_DATE)),
            mission_label=_normalize_str(_col(COL_MISSION_LABEL)),
            mission_date_raw=_parse_date_cell(_col(COL_MISSION_DATE)),
            controlled_department_raw=_normalize_str(_col(COL_CTRL_DEPT)),
            department_raw=_normalize_str(_col(COL_DEPT)),
            observations=_normalize_str(_col(COL_OBSERVATIONS)),
            anomalous_dossiers=_normalize_str(_col(COL_ANOMALIES)),
            deliverables_raw=_normalize_str(_col(COL_DELIVERABLES)),
        ))

    wb.close()
    return rows


# ── Validation (cœur, pure, sans écriture) ────────────────────────────────────


def validate_rows(rows: list[RowDraft], *, existing_refs: set[str]) -> ImportReport:
    """
    Valide les lignes parsées et retourne un ImportReport.

    Pure function : aucune écriture en base. Toutes les erreurs d'une ligne
    sont collectées (pas seulement la première) pour un rapport exhaustif (AC4).

    Args:
        rows: Lignes normalisées depuis parse_workbook.
        existing_refs: Références déjà en base (chargées par l'appelant).

    Returns:
        ImportReport avec valid_rows (dicts prêts pour Recommendation(**data)) et errors_by_row.
    """
    today = timezone.now().date()
    report = ImportReport(total=len(rows))

    # Préchargement des sources/départements actifs en dicts case-folded
    # (2 requêtes au lieu d'un SELECT par ligne — évite le N+1). setdefault +
    # tri par code reproduisent la priorité « code d'abord » de l'ancien .first().
    src_by_key: dict[str, RecommendationSource] = {}
    for s in RecommendationSource.objects.filter(is_active=True).order_by("code"):
        src_by_key.setdefault(s.code.lower(), s)
        src_by_key.setdefault(s.label.lower(), s)
    dept_by_key: dict[str, Department] = {}
    for d in Department.objects.filter(is_active=True).order_by("code"):
        dept_by_key.setdefault(d.code.lower(), d)
        dept_by_key.setdefault(d.name.lower(), d)

    # Lowercase pour comparaisons insensibles à la casse
    existing_refs_lower = {r.lower() for r in existing_refs}

    # Détection des doublons intra-fichier (bidirectionnel — AC4)
    ref_to_lines: dict[str, list[int]] = {}
    for row in rows:
        if row.reference:
            ref_to_lines.setdefault(row.reference.lower(), []).append(row.row_number)
    duplicate_lines: dict[str, list[int]] = {
        ref: lines for ref, lines in ref_to_lines.items() if len(lines) > 1
    }

    for row in rows:
        errors: list[str] = []
        data: dict = {}

        # ── Référence ──────────────────────────────────────────────────────────
        if not row.reference:
            errors.append("Référence : champ obligatoire manquant.")
        elif len(row.reference) > 50:
            errors.append(
                f"Référence : dépasse 50 caractères ({len(row.reference)} car.)."
            )
        else:
            ref_lower = row.reference.lower()
            if ref_lower in duplicate_lines:
                other = [ln for ln in duplicate_lines[ref_lower] if ln != row.row_number]
                if other:
                    errors.append(
                        f"Référence « {row.reference} » en doublon avec la ligne {other[0]}."
                    )
            if ref_lower in existing_refs_lower:
                errors.append(
                    f"Référence « {row.reference} » déjà présente en base."
                )
            data["reference"] = row.reference

        # ── Source ─────────────────────────────────────────────────────────────
        if not row.source_raw:
            errors.append("Source : champ obligatoire manquant.")
        else:
            source = src_by_key.get(row.source_raw.lower())
            if source is None:
                errors.append(
                    f"Source « {row.source_raw} » inconnue ou inactive. "
                    "Consultez l'onglet « Instructions » pour les valeurs valides."
                )
            else:
                data["source"] = source

        # ── Criticité ─────────────────────────────────────────────────────────
        if not row.priority:
            errors.append("Criticité : champ obligatoire manquant.")
        elif row.priority not in _VALID_PRIORITIES:
            errors.append(
                f"Criticité « {row.priority} » invalide. "
                "Valeurs acceptées : CRITIQUE, HAUTE, MOYENNE, FAIBLE."
            )
        else:
            data["priority"] = row.priority

        # ── Description ───────────────────────────────────────────────────────
        if not row.description:
            errors.append("Description : champ obligatoire manquant.")
        else:
            data["description"] = row.description

        # ── Date cible ────────────────────────────────────────────────────────
        if row.due_date_raw is None or row.due_date_raw == "":
            errors.append("Date cible : champ obligatoire manquant.")
        elif isinstance(row.due_date_raw, str):
            errors.append(
                f"Date cible « {row.due_date_raw} » illisible. "
                "Format attendu : AAAA-MM-JJ (ex : 2026-09-30)."
            )
        else:
            due: date = row.due_date_raw
            if due < today:
                errors.append(
                    f"Date cible {due.isoformat()} dans le passé (aujourd'hui : {today.isoformat()})."
                )
            else:
                data["due_date"] = due

        # ── Champs optionnels ─────────────────────────────────────────────────

        if row.mission_label:
            data["mission_label"] = row.mission_label

        if row.mission_date_raw is not None and row.mission_date_raw != "":
            if isinstance(row.mission_date_raw, str):
                errors.append(
                    f"Date de mission « {row.mission_date_raw} » illisible. "
                    "Format attendu : AAAA-MM-JJ."
                )
            else:
                data["mission_date"] = row.mission_date_raw

        if row.controlled_department_raw:
            dept_ctrl = dept_by_key.get(row.controlled_department_raw.lower())
            if dept_ctrl is None:
                errors.append(
                    f"Direction contrôlée « {row.controlled_department_raw} » "
                    "inexistante ou inactive."
                )
            else:
                data["controlled_department"] = dept_ctrl

        if row.department_raw:
            dept = dept_by_key.get(row.department_raw.lower())
            if dept is None:
                errors.append(
                    f"Direction concernée « {row.department_raw} » "
                    "inexistante ou inactive."
                )
            else:
                data["department"] = dept

        if row.observations:
            data["observations"] = row.observations

        if row.anomalous_dossiers:
            data["anomalous_dossiers"] = row.anomalous_dossiers

        # ── Livrables ─────────────────────────────────────────────────────────
        deliverables = [
            label.strip()
            for label in row.deliverables_raw.split(";")
            if label.strip()
        ] if row.deliverables_raw else []

        # ── Résultat ──────────────────────────────────────────────────────────
        if errors:
            report.errors_by_row[row.row_number] = errors
        else:
            report.valid_rows.append(ValidRow(data=data, deliverables_data=deliverables))

    return report


# ── Création atomique ─────────────────────────────────────────────────────────


def create_recommendations_bulk(
    *,
    rows: list[RowDraft],
    performed_by,
    uploaded_file,
    file_name: str,
    ip_address: str | None = None,
    batch: ImportBatch | None = None,
) -> ImportBatch:
    """
    Crée atomiquement toutes les recommandations d'un lot d'import.

    Re-valide DANS la transaction pour couvrir la concurrence (AC8 / Piège 7).
    Tout-ou-rien : toute erreur déclenche un rollback complet.

    Args:
        rows: Lignes parsées depuis parse_workbook.
        performed_by: Auditeur connecté.
        uploaded_file: Fichier Django pour archivage (ImportBatch.source_file).
            Ignoré si ``batch`` est fourni (le fichier est déjà archivé dessus).
        file_name: Nom original du fichier.
        ip_address: IP client pour AuditLog.
        batch: Lot déjà créé en PENDING par la vue (traitement asynchrone
            Django-Q2). Si None (appel synchrone historique), un nouveau lot
            est créé ici — comportement inchangé pour les appelants existants.

    Returns:
        ImportBatch: Le lot d'import créé (ou réutilisé) avec les recommandations liées.

    Raises:
        ValidationError: Si re-validation échoue (concurrence, date passée, etc.).
    """
    # Le fichier est écrit sur disque DANS la transaction (archivage COBAC), mais
    # l'écriture FileField n'est pas transactionnelle : on suit le FieldFile pour
    # le supprimer si la transaction rollback (sinon fichier orphelin — finding #2).
    # Un batch fourni par l'appelant (mode async) est déjà persisté avant l'appel :
    # son fichier ne doit pas être supprimé ici, il reste consultable en cas d'échec.
    batch_file = None
    created_new_batch = batch is None
    try:
        with transaction.atomic():
            # Recharger les références existantes DANS la transaction (Piège 7 —
            # concurrence AC8), requête bornée aux références du fichier (finding #4).
            current_refs = get_existing_references(rows)

            report = validate_rows(rows, existing_refs=current_refs)

            if report.errors_by_row:
                # Vérifier si c'est une collision de concurrence pour un message plus précis (AC8)
                collision_refs = []
                for errors in report.errors_by_row.values():
                    for err in errors:
                        if "déjà présente en base" in err:
                            m = re.search(r"« (.+?) »", err)
                            if m:
                                collision_refs.append(m.group(1))

                if collision_refs:
                    refs_str = ", ".join(f"« {r} »" for r in collision_refs[:3])
                    suffix = "…" if len(collision_refs) > 3 else ""
                    raise ValidationError(
                        _(
                            f"La référence {refs_str}{suffix} a été créée entre votre "
                            "prévisualisation et la confirmation. "
                            "Modifiez votre fichier et réessayez."
                        )
                    )

                error_lines = ", ".join(str(ln) for ln in sorted(report.errors_by_row.keys()))
                raise ValidationError(
                    _(
                        f"La validation a échoué sur {report.error_count} ligne(s) "
                        f"(lignes {error_lines}). Corrigez le fichier et réessayez."
                    )
                )

            # Créer le lot d'import (ou réutiliser celui fourni par l'appelant async)
            if batch is None:
                batch = ImportBatch.objects.create(
                    source_file=uploaded_file,
                    file_name=file_name,
                    created_by=performed_by,
                    recommendation_count=len(report.valid_rows),
                    total_rows=len(report.valid_rows),
                )
            else:
                batch.recommendation_count = len(report.valid_rows)
                batch.total_rows = len(report.valid_rows)
                batch.save(update_fields=["recommendation_count", "total_rows"])
            batch_file = batch.source_file  # à nettoyer si la transaction échoue

            # Boucle par ligne — PAS de bulk_create sur les recos (Piège 1 :
            # bulk_create saute full_clean() donc la règle « date non passée »
            # du Recommendation.clean() ne tournerait pas).
            created_count = 0
            for valid_row in report.valid_rows:
                row_data = dict(valid_row.data)
                # Calquer le pattern create_recommendation (services.py:58-63)
                row_data["original_due_date"] = row_data["due_date"]
                row_data["created_by"] = performed_by
                row_data["import_batch"] = batch
                # import_tag reste null (Piège 2 — lèverait l'exemption de date passée)

                recommendation = Recommendation(**row_data)
                recommendation.full_clean(exclude=["status"])
                try:
                    recommendation.save()
                except IntegrityError:
                    # Course non visible à la re-validation : la référence a été
                    # commitée par un import concurrent entre le re-check et ce save.
                    # Convertir en ValidationError pour le message gracieux AC8
                    # (la vue ne capture que DjangoValidationError — finding #1).
                    raise ValidationError(
                        _(
                            f"La référence « {recommendation.reference} » a été créée "
                            "entre votre prévisualisation et la confirmation. "
                            "Modifiez votre fichier et réessayez."
                        )
                    )

                # Livrables — bulk_create OK après save du parent (seul cas autorisé)
                deliverables = [
                    Deliverable(recommendation=recommendation, label=label, order=order)
                    for order, label in enumerate(valid_row.deliverables_data)
                ]
                if deliverables:
                    Deliverable.objects.bulk_create(deliverables)

                # AuditLog CREATE par reco (traçabilité individuelle)
                AuditLog.objects.create(
                    action=AuditLog.Action.CREATE,
                    user=performed_by,
                    content_type="Recommendation",
                    object_id=recommendation.pk,
                    changes={
                        "reference": recommendation.reference,
                        "source": recommendation.source.code if recommendation.source else None,
                        "priority": recommendation.priority,
                        "status": recommendation.status,
                        "deliverables_count": len(deliverables),
                        "import_batch": str(batch.pk),
                    },
                    description=(
                        f"Création de la recommandation {recommendation.reference} "
                        f"(lot d'import {batch.pk}) par {performed_by.username}"
                    ),
                    ip_address=ip_address,
                )
                created_count += 1
                # Progression pour le polling de suivi (mode async) : mise à jour tous
                # les 25 lignes, pas à chaque ligne, pour ne pas multiplier les UPDATE.
                # updated_at est touché à chaque fois pour signaler au reconciliateur
                # de batches bloqués (reconcile_stale_processing_batches) qu'un import
                # long est toujours actif, pas planté.
                if created_count % 25 == 0:
                    ImportBatch.objects.filter(pk=batch.pk).update(
                        processed_rows=created_count, updated_at=timezone.now(),
                    )

            # Valeur finale explicite : le compteur périodique ci-dessus saute la
            # dernière tranche si created_count n'est pas multiple de 25 (import_pending.html
            # afficherait sinon une progression figée sous 100 % jusqu'au statut DONE).
            ImportBatch.objects.filter(pk=batch.pk).update(
                processed_rows=created_count, updated_at=timezone.now(),
            )

            # AuditLog IMPORT de batch — une seule entrée pour le lot (JSON-safe, Piège 6)
            AuditLog.objects.create(
                action=AuditLog.Action.IMPORT,
                user=performed_by,
                content_type="ImportBatch",
                object_id=batch.pk,
                changes={
                    "file": file_name,
                    "count": created_count,
                    "batch": str(batch.pk),
                },
                description=(
                    f"Import Excel de {created_count} recommandation(s) "
                    f"par {performed_by.username}"
                ),
                ip_address=ip_address,
            )

        return batch
    except Exception:
        # Transaction rollback → supprimer le fichier physique déjà écrit pour
        # ne pas laisser d'orphelin dans media/imports/ (finding #2). En mode
        # async (batch fourni par l'appelant), le fichier reste archivé pour
        # diagnostic — c'est run_recommendations_import_task qui gère l'échec.
        if batch_file and created_new_batch:
            batch_file.delete(save=False)
        raise


def _notify_import_result(batch: ImportBatch, performed_by, *, success: bool, error: str = "") -> None:
    """Notifie l'auditeur qui a lancé l'import de l'issue du traitement asynchrone."""
    from apps.notifications.services import emit_notification
    from apps.notifications.models import Notification

    if success:
        emit_notification(
            recipient=performed_by,
            notification_type=Notification.Type.IMPORT_COMPLETED,
            title=f"Import « {batch.file_name} » terminé",
            idempotency_key=f"IMPORT_COMPLETED:{batch.pk}",
            body=f"{batch.recommendation_count} recommandation(s) créée(s).",
            url=f"/workflow/recommandations/import/status/{batch.pk}/",
        )
    else:
        emit_notification(
            recipient=performed_by,
            notification_type=Notification.Type.IMPORT_FAILED,
            title=f"Échec de l'import « {batch.file_name} »",
            idempotency_key=f"IMPORT_FAILED:{batch.pk}",
            body=error[:500],
            url=f"/workflow/recommandations/import/status/{batch.pk}/",
            is_urgent=True,
        )


def run_recommendations_import_task(
    *, batch_id: str, performed_by_id: int, ip_address: str | None = None
) -> str:
    """
    Point d'entrée Django-Q2 pour l'import Excel massif.

    Relit le fichier archivé dans ``ImportBatch.source_file`` (sauvegardé par
    la vue avant l'enqueue), re-parse et délègue à ``create_recommendations_bulk``
    (logique métier inchangée, couverte par les tests synchrones existants).

    Idempotence : si le batch n'est plus PENDING (déjà DONE/PROCESSING/CANCELLED
    par un run antérieur ou une annulation), sortie immédiate — protège contre
    un retry Django-Q2 après un succès déclaré tardivement.

    Le nom de cette fonction est sérialisé tel quel par Django-Q2 dans la table
    des tâches : ne pas la renommer sans migrer les tâches en vol.
    """
    from django.contrib.auth import get_user_model

    with transaction.atomic():
        batch = ImportBatch.objects.select_for_update().get(pk=batch_id)
        if batch.status != ImportBatch.Status.PENDING:
            return f"skip: batch {batch_id} already {batch.status}"
        batch.status = ImportBatch.Status.PROCESSING
        batch.save(update_fields=["status"])

    performed_by = get_user_model().objects.get(pk=performed_by_id)

    try:
        with batch.source_file.open("rb") as f:
            rows = parse_workbook(f)
            create_recommendations_bulk(
                rows=rows,
                performed_by=performed_by,
                uploaded_file=None,
                file_name=batch.file_name,
                ip_address=ip_address,
                batch=batch,
            )
        batch.refresh_from_db()
        batch.status = ImportBatch.Status.DONE
        batch.save(update_fields=["status"])
        _notify_import_result(batch, performed_by, success=True)
        return "done"
    except Exception as exc:
        error_message = str(exc)[:2000]
        ImportBatch.objects.filter(pk=batch_id).update(
            status=ImportBatch.Status.FAILED,
            error_message=error_message,
        )
        # La transaction de create_recommendations_bulk a fait un rollback complet —
        # aucune trace de l'échec n'existe dans l'AuditLog. On la crée ici, hors
        # transaction annulée, pour que l'échec reste visible dans le journal d'audit.
        AuditLog.objects.create(
            action=AuditLog.Action.SYSTEM,
            user=performed_by,
            content_type="ImportBatch",
            object_id=batch.pk,
            changes={"file": batch.file_name, "error": error_message},
            description=f"Échec de l'import Excel « {batch.file_name} » : {error_message}",
            ip_address=ip_address,
        )
        _notify_import_result(batch, performed_by, success=False, error=error_message)
        raise
