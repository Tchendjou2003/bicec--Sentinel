"""
Workflow App — Tests Import Historique (Story 6.8)

Couvre :
  - parse_historical_workbook : format, plafond, onglet manquant
  - validate_historical_rows  : dates passées OK, clôture future KO, cohérence dates,
                                 doublons, références existantes, champs obligatoires
  - parse_zip_members         : Zip Bomb, Zip Slip, fichier invalide
  - _match_zip_to_refs        : convention {reference}_*.ext
  - create_historical_recommendations : atomicité, HMAC, AuditLog, preuves, created_at
  - _ensure_not_closed        : protection read-only sur IMPORTED
  - Vues                      : RBAC, HTMX preview, confirm
"""
from __future__ import annotations

import io
import os
import tempfile
import zipfile
from datetime import date, timedelta
from unittest import mock

import openpyxl
from django.core.exceptions import ValidationError
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, override_settings
from django.urls import reverse
from django.utils import timezone

from apps.audit.models import AuditLog, HmacSeal
from apps.users.models import Department, OrgUnitType, User
from apps.workflow.import_historical import (
    DATA_SHEET_NAME,
    HEADERS,
    MAX_ROWS,
    HistoricalImportReport,
    HistoricalRow,
    ValidHistoricalRow,
    ZipFileEntry,
    _match_zip_to_refs,
    build_historical_import_template,
    create_historical_recommendations,
    parse_historical_workbook,
    parse_zip_members,
    run_historical_import_task,
    validate_historical_rows,
)
from apps.workflow.models import (
    EvidenceFile,
    EvidenceSubmission,
    ImportBatch,
    Recommendation,
    RecommendationSource,
)


# ── Fixtures partagées ────────────────────────────────────────────────────────


class HistoricalImportMixin:
    """Données de test communes."""

    @classmethod
    def setUpTestData(cls):
        cls.org_type, _ = OrgUnitType.objects.get_or_create(
            code="DIRECTION", defaults={"name": "Direction", "level": 1},
        )
        cls.dept = Department.objects.create(
            name="Direction Opérations", code="DOP", type=cls.org_type,
        )
        cls.source, _ = RecommendationSource.objects.get_or_create(
            code="COBAC", defaults={"label": "COBAC", "is_external": True},
        )
        cls.audit_user = User.objects.create_user(
            username="audit_hist", password="TestPass123!", role=User.Role.AUDIT,
        )
        cls.dm_user = User.objects.create_user(
            username="dm_hist", password="TestPass123!", role=User.Role.DM,
        )

    # Helpers

    @staticmethod
    def _make_excel(rows: list[list], sheet: str = DATA_SHEET_NAME) -> io.BytesIO:
        """Crée un .xlsx en mémoire avec les lignes fournies."""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = sheet
        ws.append(HEADERS)
        for row in rows:
            ws.append(row)
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return buf

    @staticmethod
    def _make_zip(files: dict[str, bytes]) -> io.BytesIO:
        """Crée un ZIP en mémoire {nom_fichier: contenu}."""
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
            for name, data in files.items():
                zf.writestr(name, data)
        buf.seek(0)
        return buf

    def _make_valid_row_data(
        self,
        reference: str = "RECO-2022-001",
        closed_at: date | None = None,
        created_at: date | None = None,
    ) -> list:
        """Retourne une ligne Excel valide."""
        today = date.today()
        return [
            reference,                           # Référence
            "COBAC",                             # Source
            "HAUTE",                             # Criticité
            "Description de test historique",   # Description
            (today - timedelta(days=365)).isoformat(),  # Date cible originale
            (closed_at or today - timedelta(days=30)).isoformat(),   # Date de clôture
            (created_at or today - timedelta(days=400)).isoformat(), # Date de création
            "Mission 2022",                      # Titre (mission)
            "",                                  # Direction contrôlée
            "",                                  # Direction concernée
            "RAS",                               # Observations
        ]


# ── Tests : parse_historical_workbook ─────────────────────────────────────────


class ParseHistoricalWorkbookTests(HistoricalImportMixin, TestCase):

    @mock.patch("apps.workflow.import_historical.validate_magic_bytes")
    @mock.patch("apps.workflow.import_historical.validate_file_size")
    def test_parse_valid_workbook(self, mock_size, mock_magic):
        buf = self._make_excel([self._make_valid_row_data()])
        buf.name = "test.xlsx"
        rows = parse_historical_workbook(buf)
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0].reference, "RECO-2022-001")

    @mock.patch("apps.workflow.import_historical.validate_magic_bytes")
    @mock.patch("apps.workflow.import_historical.validate_file_size")
    def test_parse_empty_rows_skipped(self, mock_size, mock_magic):
        buf = self._make_excel([])
        buf.name = "test.xlsx"
        rows = parse_historical_workbook(buf)
        self.assertEqual(rows, [])

    @mock.patch("apps.workflow.import_historical.validate_magic_bytes")
    @mock.patch("apps.workflow.import_historical.validate_file_size")
    def test_wrong_sheet_name_raises(self, mock_size, mock_magic):
        buf = self._make_excel([self._make_valid_row_data()], sheet="WrongSheet")
        buf.name = "test.xlsx"
        with self.assertRaises(ValidationError) as ctx:
            parse_historical_workbook(buf)
        self.assertIn("Données", str(ctx.exception))

    @mock.patch("apps.workflow.import_historical.validate_magic_bytes")
    @mock.patch("apps.workflow.import_historical.validate_file_size")
    def test_max_rows_exceeded_raises(self, mock_size, mock_magic):
        rows_data = [self._make_valid_row_data(f"RECO-{i:04d}") for i in range(MAX_ROWS + 1)]
        buf = self._make_excel(rows_data)
        buf.name = "test.xlsx"
        with self.assertRaises(ValidationError) as ctx:
            parse_historical_workbook(buf)
        self.assertIn(str(MAX_ROWS), str(ctx.exception))


# ── Tests : validate_historical_rows ─────────────────────────────────────────


class ValidateHistoricalRowsTests(HistoricalImportMixin, TestCase):

    def _row(self, **kwargs) -> HistoricalRow:
        today = date.today()
        defaults = dict(
            row_number=2,
            reference="RECO-2022-001",
            source_raw="COBAC",
            priority="HAUTE",
            description="Desc test",
            due_date_raw=today - timedelta(days=365),
            closed_at_raw=today - timedelta(days=30),
            created_at_original_raw=today - timedelta(days=400),
            mission_label="Mission 2022",
            controlled_department_raw="",
            department_raw="",
            observations="",
        )
        defaults.update(kwargs)
        return HistoricalRow(**defaults)

    def test_valid_row_produces_valid_count(self):
        report = validate_historical_rows([self._row()])
        self.assertEqual(report.valid_count, 1)
        self.assertEqual(report.error_count, 0)

    def test_past_due_date_is_allowed(self):
        """Les dates passées sont autorisées pour les recos historiques."""
        today = date.today()
        row = self._row(due_date_raw=today - timedelta(days=2000))
        report = validate_historical_rows([row])
        self.assertEqual(report.valid_count, 1)

    def test_future_closed_at_is_rejected(self):
        """Une date de clôture dans le futur est invalide."""
        row = self._row(closed_at_raw=date.today() + timedelta(days=10))
        report = validate_historical_rows([row])
        self.assertEqual(report.error_count, 1)
        errors = list(report.errors_by_row.values())[0]
        self.assertTrue(any("futur" in e for e in errors))

    def test_created_at_after_closed_at_is_rejected(self):
        """La date de création doit être antérieure à la date de clôture."""
        today = date.today()
        row = self._row(
            closed_at_raw=today - timedelta(days=100),
            created_at_original_raw=today - timedelta(days=50),  # après clôture
        )
        report = validate_historical_rows([row])
        self.assertEqual(report.error_count, 1)
        errors = list(report.errors_by_row.values())[0]
        self.assertTrue(any("postérieure" in e for e in errors))

    def test_missing_reference_produces_error(self):
        row = self._row(reference="")
        report = validate_historical_rows([row])
        self.assertEqual(report.error_count, 1)
        errors = list(report.errors_by_row.values())[0]
        self.assertTrue(any("Référence" in e for e in errors))

    def test_invalid_priority_produces_error(self):
        row = self._row(priority="UNKNOWN")
        report = validate_historical_rows([row])
        self.assertEqual(report.error_count, 1)

    def test_unknown_source_produces_error(self):
        row = self._row(source_raw="SOURCE_INEXISTANTE")
        report = validate_historical_rows([row])
        self.assertEqual(report.error_count, 1)

    def test_intra_file_duplicate_produces_error(self):
        rows = [self._row(row_number=2), self._row(row_number=3)]
        report = validate_historical_rows(rows)
        self.assertEqual(report.error_count, 2)
        all_errors = [e for errs in report.errors_by_row.values() for e in errs]
        self.assertTrue(any("doublon" in e for e in all_errors))

    def test_existing_reference_produces_error(self):
        Recommendation.objects.create(
            reference="RECO-EXIST",
            description="Existante",
            source=self.source,
            priority="HAUTE",
            due_date=date.today() + timedelta(days=30),
            original_due_date=date.today() + timedelta(days=30),
            created_by=self.audit_user,
        )
        row = self._row(reference="RECO-EXIST")
        report = validate_historical_rows([row])
        self.assertEqual(report.error_count, 1)
        errors = list(report.errors_by_row.values())[0]
        self.assertTrue(any("déjà présente en base" in e for e in errors))

    def test_zip_matching_in_report(self):
        row = self._row()
        zip_entries = [
            ZipFileEntry(name="RECO-2022-001_rapport.pdf", data=b"PDF"),
            ZipFileEntry(name="AUTRE_fichier.pdf", data=b"PDF"),
        ]
        report = validate_historical_rows([row], zip_entries)
        self.assertEqual(report.zip_total, 2)
        self.assertEqual(report.zip_matched, 1)
        self.assertEqual(report.zip_unmatched_count, 1)
        self.assertIn("AUTRE_fichier.pdf", report.zip_unmatched)


# ── Tests : _match_zip_to_refs ────────────────────────────────────────────────


class MatchZipToRefsTests(TestCase):

    def test_exact_prefix_match(self):
        entries = [ZipFileEntry("RECO-001_rapport.pdf", b"")]
        matched, unmatched = _match_zip_to_refs(entries, ["RECO-001"])
        self.assertEqual(len(matched["RECO-001"]), 1)
        self.assertEqual(unmatched, [])

    def test_case_insensitive_match(self):
        entries = [ZipFileEntry("reco-001_rapport.pdf", b"")]
        matched, unmatched = _match_zip_to_refs(entries, ["RECO-001"])
        self.assertEqual(len(matched["RECO-001"]), 1)

    def test_no_separator_does_not_match(self):
        """Un fichier sans underscore après la référence ne doit pas matcher."""
        entries = [ZipFileEntry("RECO-001rapport.pdf", b"")]
        _, unmatched = _match_zip_to_refs(entries, ["RECO-001"])
        self.assertIn("RECO-001rapport.pdf", unmatched)

    def test_multiple_files_same_ref(self):
        entries = [
            ZipFileEntry("RECO-001_a.pdf", b""),
            ZipFileEntry("RECO-001_b.pdf", b""),
        ]
        matched, unmatched = _match_zip_to_refs(entries, ["RECO-001"])
        self.assertEqual(len(matched["RECO-001"]), 2)
        self.assertEqual(unmatched, [])

    def test_folder_based_match(self):
        """Fichier dans un dossier nommé comme la référence → associé sans renommage."""
        entries = [ZipFileEntry("rapport.pdf", b"", ref_hint="RECO-001")]
        matched, unmatched = _match_zip_to_refs(entries, ["RECO-001"])
        self.assertEqual(len(matched["RECO-001"]), 1)
        self.assertEqual(unmatched, [])

    def test_folder_based_multi_files(self):
        """Plusieurs fichiers dans le même dossier → tous associés à la même reco."""
        entries = [
            ZipFileEntry("rapport.pdf", b"", ref_hint="RECO-001"),
            ZipFileEntry("pv.pdf", b"", ref_hint="RECO-001"),
        ]
        matched, unmatched = _match_zip_to_refs(entries, ["RECO-001"])
        self.assertEqual(len(matched["RECO-001"]), 2)
        self.assertEqual(unmatched, [])

    def test_prefix_fallback_still_works(self):
        """Un fichier à plat avec préfixe (ancienne convention) reste matché."""
        entries = [ZipFileEntry("RECO-002_contrat.pdf", b"", ref_hint="")]
        matched, unmatched = _match_zip_to_refs(entries, ["RECO-002"])
        self.assertEqual(len(matched["RECO-002"]), 1)
        self.assertEqual(unmatched, [])

    def test_folder_takes_priority_and_prefix_fallback_coexist(self):
        """Les deux méthodes coexistent dans le même ZIP sans interférence."""
        entries = [
            ZipFileEntry("rapport.pdf", b"", ref_hint="RECO-001"),    # méthode dossier
            ZipFileEntry("RECO-002_note.pdf", b"", ref_hint=""),       # méthode préfixe
        ]
        matched, unmatched = _match_zip_to_refs(entries, ["RECO-001", "RECO-002"])
        self.assertEqual(len(matched["RECO-001"]), 1)
        self.assertEqual(len(matched["RECO-002"]), 1)
        self.assertEqual(unmatched, [])

    def test_folder_case_insensitive(self):
        """Le matching par dossier est insensible à la casse."""
        entries = [ZipFileEntry("doc.pdf", b"", ref_hint="reco-001")]
        matched, unmatched = _match_zip_to_refs(entries, ["RECO-001"])
        self.assertEqual(len(matched["RECO-001"]), 1)
        self.assertEqual(unmatched, [])


# ── Tests : parse_zip_members ─────────────────────────────────────────────────


class ParseZipMembersTests(TestCase):

    @mock.patch("apps.workflow.import_historical.validate_magic_bytes")
    @mock.patch("apps.workflow.import_historical.validate_file_size")
    def test_valid_zip_returns_entries(self, mock_size, mock_magic):
        buf = self._make_zip({"RECO-001_rapport.pdf": b"PDF content"})
        entries = parse_zip_members(buf)
        self.assertEqual(len(entries), 1)
        self.assertEqual(entries[0].name, "RECO-001_rapport.pdf")

    @mock.patch("apps.workflow.import_historical.validate_file_size")
    def test_not_a_zip_raises(self, mock_size):
        buf = io.BytesIO(b"This is not a ZIP file at all.")
        buf.name = "notazip.zip"
        with self.assertRaises(ValidationError) as ctx:
            parse_zip_members(buf)
        self.assertIn("valide", str(ctx.exception))

    @mock.patch("apps.workflow.import_historical.validate_magic_bytes")
    @mock.patch("apps.workflow.import_historical.validate_file_size")
    def test_zip_slip_protection(self, mock_size, mock_magic):
        """Les noms de fichier traversant des répertoires parents sont neutralisés."""
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w") as zf:
            zf.writestr("../../../evil.pdf", b"evil content")
        buf.seek(0)
        entries = parse_zip_members(buf)
        # Le nom de fichier safe ne doit pas contenir de traversée
        for entry in entries:
            self.assertNotIn("..", entry.name)
            self.assertNotIn("/", entry.name)
            # ref_hint ne doit pas être ".." — un dossier parent traversant
            # n'est jamais une référence valide donc ne serait pas matché,
            # mais on documente explicitement la propriété
            self.assertNotEqual(entry.ref_hint, "..")

    @mock.patch("apps.workflow.import_historical.validate_magic_bytes")
    @mock.patch("apps.workflow.import_historical.validate_file_size")
    def test_zip_bomb_rejected(self, mock_size, mock_magic):
        """Un ZIP dont le total décompressé > 100 Mo est rejeté."""
        buf = io.BytesIO()
        large_data = b"0" * (101 * 1024 * 1024)  # 101 Mo
        with zipfile.ZipFile(buf, "w", zipfile.ZIP_STORED) as zf:
            zf.writestr("big.pdf", large_data)
        buf.seek(0)
        with self.assertRaises(ValidationError) as ctx:
            parse_zip_members(buf)
        self.assertIn("100 Mo", str(ctx.exception))

    @staticmethod
    def _make_zip(files: dict[str, bytes]) -> io.BytesIO:
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w") as zf:
            for name, data in files.items():
                zf.writestr(name, data)
        buf.seek(0)
        return buf


# ── Tests : create_historical_recommendations ─────────────────────────────────


@override_settings(
    MEDIA_ROOT=tempfile.mkdtemp(),
    HMAC_SECRET_KEY="test-secret-key-sentinel",
)
class CreateHistoricalRecommendationsTests(HistoricalImportMixin, TestCase):

    def _valid_rows(
        self, reference: str = "RECO-HIST-001"
    ) -> list[HistoricalRow]:
        today = date.today()
        return [
            HistoricalRow(
                row_number=2,
                reference=reference,
                source_raw="COBAC",
                priority="HAUTE",
                description="Reco historique de test",
                due_date_raw=today - timedelta(days=365),
                closed_at_raw=today - timedelta(days=30),
                created_at_original_raw=today - timedelta(days=400),
                mission_label="Mission 2022",
                controlled_department_raw="",
                department_raw="",
                observations="",
            )
        ]

    def _make_excel_file(self, rows: list[HistoricalRow]) -> SimpleUploadedFile:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = DATA_SHEET_NAME
        ws.append(HEADERS)
        today = date.today()
        for row in rows:
            ws.append([
                row.reference,
                row.source_raw,
                row.priority,
                row.description,
                row.due_date_raw.isoformat() if row.due_date_raw else "",
                row.closed_at_raw.isoformat() if row.closed_at_raw else "",
                row.created_at_original_raw.isoformat() if row.created_at_original_raw else "",
                row.mission_label,
                row.controlled_department_raw,
                row.department_raw,
                row.observations,
            ])
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return SimpleUploadedFile("test_historique.xlsx", buf.read(),
                                  content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

    @mock.patch("apps.workflow.import_historical.validate_magic_bytes")
    @mock.patch("apps.workflow.import_historical.validate_file_size")
    def test_creates_recommendation_with_correct_status(self, mock_size, mock_magic):
        rows = self._valid_rows()
        excel = self._make_excel_file(rows)
        batch = create_historical_recommendations(
            rows=rows,
            zip_entries=None,
            performed_by=self.audit_user,
            uploaded_file=excel,
            file_name="test.xlsx",
        )
        reco = Recommendation.all_objects.get(reference="RECO-HIST-001")
        self.assertEqual(reco.status, Recommendation.Status.CLOSED_RESOLVED)
        self.assertEqual(reco.import_tag, "IMPORTED")
        self.assertEqual(str(reco.import_batch.pk), str(batch.pk))

    @mock.patch("apps.workflow.import_historical.validate_magic_bytes")
    @mock.patch("apps.workflow.import_historical.validate_file_size")
    def test_created_at_overridden_from_excel(self, mock_size, mock_magic):
        """created_at doit refléter la date historique et non timezone.now()."""
        rows = self._valid_rows()
        excel = self._make_excel_file(rows)
        create_historical_recommendations(
            rows=rows,
            zip_entries=None,
            performed_by=self.audit_user,
            uploaded_file=excel,
            file_name="test.xlsx",
        )
        reco = Recommendation.all_objects.get(reference="RECO-HIST-001")
        expected_date = date.today() - timedelta(days=400)
        self.assertEqual(reco.created_at.date(), expected_date)

    @mock.patch("apps.workflow.import_historical.validate_magic_bytes")
    @mock.patch("apps.workflow.import_historical.validate_file_size")
    def test_hmac_seal_generated(self, mock_size, mock_magic):
        rows = self._valid_rows()
        excel = self._make_excel_file(rows)
        create_historical_recommendations(
            rows=rows,
            zip_entries=None,
            performed_by=self.audit_user,
            uploaded_file=excel,
            file_name="test.xlsx",
        )
        reco = Recommendation.all_objects.get(reference="RECO-HIST-001")
        self.assertTrue(hasattr(reco, "hmac_seal"))
        seal = reco.hmac_seal
        self.assertIsNotNone(seal)
        self.assertEqual(seal.sealed_by, self.audit_user)

    @mock.patch("apps.workflow.import_historical.validate_magic_bytes")
    @mock.patch("apps.workflow.import_historical.validate_file_size")
    def test_audit_log_created(self, mock_size, mock_magic):
        rows = self._valid_rows()
        excel = self._make_excel_file(rows)
        create_historical_recommendations(
            rows=rows,
            zip_entries=None,
            performed_by=self.audit_user,
            uploaded_file=excel,
            file_name="test.xlsx",
            provenance="Archives papier numérisées",
        )
        log = AuditLog.objects.filter(
            action=AuditLog.Action.IMPORT,
            user=self.audit_user,
        ).first()
        self.assertIsNotNone(log)
        self.assertIn("Archives papier", log.description)

    @mock.patch("apps.workflow.import_historical.validate_magic_bytes")
    @mock.patch("apps.workflow.import_historical.validate_file_size")
    def test_zip_evidence_created(self, mock_size, mock_magic):
        """Les preuves ZIP sont créées comme EvidenceFile avec status ACCEPTED."""
        rows = self._valid_rows("RECO-HIST-002")
        zip_entries = [
            ZipFileEntry(
                name="RECO-HIST-002_preuve.pdf",
                data=b"%PDF-test content",
            )
        ]
        excel = self._make_excel_file(rows)
        create_historical_recommendations(
            rows=rows,
            zip_entries=zip_entries,
            performed_by=self.audit_user,
            uploaded_file=excel,
            file_name="test.xlsx",
        )
        reco = Recommendation.all_objects.get(reference="RECO-HIST-002")
        sub = reco.evidence_submissions.filter(
            status=EvidenceSubmission.SubmissionStatus.ACCEPTED
        ).first()
        self.assertIsNotNone(sub)
        self.assertEqual(sub.files.count(), 1)
        ev = sub.files.first()
        self.assertEqual(ev.original_filename, "RECO-HIST-002_preuve.pdf")

    @mock.patch("apps.workflow.import_historical.validate_magic_bytes")
    @mock.patch("apps.workflow.import_historical.validate_file_size")
    def test_atomic_rollback_on_duplicate(self, mock_size, mock_magic):
        """Toute erreur provoque un rollback complet (atomicité)."""
        # Créer une reco existante pour provoquer une erreur de validation
        Recommendation.objects.create(
            reference="RECO-HIST-DUP",
            description="Existante",
            source=self.source,
            priority="HAUTE",
            due_date=date.today() + timedelta(days=30),
            original_due_date=date.today() + timedelta(days=30),
            created_by=self.audit_user,
        )
        rows = self._valid_rows("RECO-HIST-DUP")
        excel = self._make_excel_file(rows)
        initial_count = Recommendation.all_objects.count()
        with self.assertRaises(ValidationError):
            create_historical_recommendations(
                rows=rows,
                zip_entries=None,
                performed_by=self.audit_user,
                uploaded_file=excel,
                file_name="test.xlsx",
            )
        # Aucune nouvelle reco créée
        self.assertEqual(Recommendation.all_objects.count(), initial_count)


# ── Tests : run_historical_import_task (Django-Q2 async) ──────────────────────


class RunHistoricalImportTaskTests(HistoricalImportMixin, TestCase):

    def _make_pending_batch(self, rows: list[HistoricalRow], file_name="hist_async.xlsx") -> ImportBatch:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = DATA_SHEET_NAME
        ws.append(HEADERS)
        for row in rows:
            ws.append([
                row.reference, row.source_raw, row.priority, row.description,
                row.due_date_raw.isoformat() if row.due_date_raw else "",
                row.closed_at_raw.isoformat() if row.closed_at_raw else "",
                row.created_at_original_raw.isoformat() if row.created_at_original_raw else "",
                row.mission_label, row.controlled_department_raw,
                row.department_raw, row.observations,
            ])
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return ImportBatch.objects.create(
            source_file=SimpleUploadedFile(
                file_name, buf.read(),
                content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            ),
            file_name=file_name,
            created_by=self.audit_user,
            status=ImportBatch.Status.PENDING,
            kind=ImportBatch.Kind.HISTORICAL,
        )

    @mock.patch("apps.workflow.import_historical.validate_magic_bytes")
    @mock.patch("apps.workflow.import_historical.validate_file_size")
    def test_traitement_reussi_cree_la_reco_et_le_sceau(self, mock_size, mock_magic):
        rows = [HistoricalRow(
            row_number=2, reference="RECO-ASYNC-001", source_raw="COBAC", priority="HAUTE",
            description="Reco async", due_date_raw=date.today() - timedelta(days=365),
            closed_at_raw=date.today() - timedelta(days=30),
            created_at_original_raw=date.today() - timedelta(days=400),
            mission_label="Mission 2022", controlled_department_raw="",
            department_raw="", observations="",
        )]
        batch = self._make_pending_batch(rows)
        result = run_historical_import_task(
            batch_id=str(batch.pk), performed_by_id=self.audit_user.pk,
        )
        batch.refresh_from_db()
        self.assertEqual(result, "done")
        self.assertEqual(batch.status, ImportBatch.Status.DONE)
        reco = Recommendation.all_objects.get(reference="RECO-ASYNC-001")
        self.assertEqual(reco.status, Recommendation.Status.CLOSED_RESOLVED)
        self.assertTrue(HmacSeal.objects.filter(recommendation=reco).exists())

    @mock.patch("apps.workflow.import_historical.validate_magic_bytes")
    @mock.patch("apps.workflow.import_historical.validate_file_size")
    def test_retry_apres_succes_ne_duplique_pas_le_sceau(self, mock_size, mock_magic):
        """Idempotence : un retry Django-Q2 après un DONE ne recrée ni reco ni sceau."""
        rows = [HistoricalRow(
            row_number=2, reference="RECO-ASYNC-RETRY", source_raw="COBAC", priority="HAUTE",
            description="Reco retry", due_date_raw=date.today() - timedelta(days=365),
            closed_at_raw=date.today() - timedelta(days=30),
            created_at_original_raw=date.today() - timedelta(days=400),
            mission_label="Mission 2022", controlled_department_raw="",
            department_raw="", observations="",
        )]
        batch = self._make_pending_batch(rows)
        run_historical_import_task(batch_id=str(batch.pk), performed_by_id=self.audit_user.pk)
        reco = Recommendation.all_objects.get(reference="RECO-ASYNC-RETRY")

        result = run_historical_import_task(
            batch_id=str(batch.pk), performed_by_id=self.audit_user.pk,
        )
        self.assertIn("skip", result)
        self.assertEqual(
            Recommendation.all_objects.filter(reference="RECO-ASYNC-RETRY").count(), 1,
        )
        self.assertEqual(HmacSeal.objects.filter(recommendation=reco).count(), 1)

    @mock.patch("apps.workflow.import_historical.validate_magic_bytes")
    @mock.patch("apps.workflow.import_historical.validate_file_size")
    def test_zip_temporaire_supprime_apres_traitement(self, mock_size, mock_magic):
        from django.core.files.storage import default_storage

        rows = [HistoricalRow(
            row_number=2, reference="RECO-ASYNC-ZIP", source_raw="COBAC", priority="HAUTE",
            description="Reco zip", due_date_raw=date.today() - timedelta(days=365),
            closed_at_raw=date.today() - timedelta(days=30),
            created_at_original_raw=date.today() - timedelta(days=400),
            mission_label="Mission 2022", controlled_department_raw="",
            department_raw="", observations="",
        )]
        batch = self._make_pending_batch(rows)
        zip_buf = self._make_zip({"RECO-ASYNC-ZIP_preuve.pdf": b"%PDF-test"})
        zip_path = default_storage.save(f"imports/zip/{batch.pk}.zip", zip_buf)

        run_historical_import_task(
            batch_id=str(batch.pk), performed_by_id=self.audit_user.pk, zip_storage_path=zip_path,
        )
        self.assertFalse(default_storage.exists(zip_path))

    def test_echec_marque_failed_avec_audit_log_hors_transaction(self):
        rows = [HistoricalRow(
            row_number=2, reference="RECO-ASYNC-FAIL", source_raw="COBAC", priority="HAUTE",
            description="Reco fail", due_date_raw=date.today() - timedelta(days=365),
            closed_at_raw=date.today() - timedelta(days=30),
            created_at_original_raw=date.today() - timedelta(days=400),
            mission_label="Mission 2022", controlled_department_raw="",
            department_raw="", observations="",
        )]
        batch = self._make_pending_batch(rows)
        with mock.patch(
            "apps.workflow.import_historical.parse_historical_workbook",
            side_effect=ValidationError("Fichier corrompu"),
        ):
            with self.assertRaises(ValidationError):
                run_historical_import_task(
                    batch_id=str(batch.pk), performed_by_id=self.audit_user.pk,
                )
        batch.refresh_from_db()
        self.assertEqual(batch.status, ImportBatch.Status.FAILED)
        self.assertTrue(batch.error_message)
        self.assertTrue(
            AuditLog.objects.filter(
                action=AuditLog.Action.SYSTEM,
                content_type="ImportBatch",
                object_id=batch.pk,
            ).exists()
        )


# ── Tests : protection read-only (_ensure_not_closed) ─────────────────────────


class ImportedReadOnlyTests(HistoricalImportMixin, TestCase):

    @override_settings(HMAC_SECRET_KEY="test-secret-key-sentinel")
    def test_imported_reco_is_blocked_by_ensure_not_closed(self):
        """Une reco IMPORTED doit être rejetée par _ensure_not_closed."""
        from apps.workflow.views import _ensure_not_closed
        reco = Recommendation.objects.create(
            reference="RECO-BLOCK-TEST",
            description="Importée",
            source=self.source,
            priority="HAUTE",
            due_date=date.today() - timedelta(days=30),
            original_due_date=date.today() - timedelta(days=30),
            import_tag="IMPORTED",
            import_batch=ImportBatch.objects.create(
                file_name="t.xlsx",
                created_by=self.audit_user,
                recommendation_count=1,
            ),
            status=Recommendation.Status.CLOSED_RESOLVED,
            created_by=self.audit_user,
            closed_by=self.audit_user,
            closed_at=timezone.now(),
        )
        with self.assertRaises(ValueError) as ctx:
            _ensure_not_closed(reco)
        self.assertIn("historique importé", str(ctx.exception))


# ── Tests : vues RBAC ─────────────────────────────────────────────────────────


class HistoricalImportViewRbacTests(HistoricalImportMixin, TestCase):

    def test_historical_import_requires_login(self):
        url = reverse("workflow:historical-import")
        response = self.client.get(url)
        self.assertNotEqual(response.status_code, 200)

    def test_dm_cannot_access_historical_import(self):
        self.client.force_login(self.dm_user)
        url = reverse("workflow:historical-import")
        response = self.client.get(url)
        self.assertIn(response.status_code, [302, 403])

    def test_audit_can_access_historical_import(self):
        self.client.force_login(self.audit_user)
        url = reverse("workflow:historical-import")
        response = self.client.get(url)
        self.assertEqual(response.status_code, 200)
        self.assertTemplateUsed(response, "workflow/historical_import.html")

    def test_template_download_requires_audit(self):
        url = reverse("workflow:historical-import-template")
        self.client.force_login(self.dm_user)
        response = self.client.get(url)
        self.assertIn(response.status_code, [302, 403])

    def test_template_download_returns_xlsx(self):
        self.client.force_login(self.audit_user)
        url = reverse("workflow:historical-import-template")
        response = self.client.get(url)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

    @mock.patch("apps.workflow.views.parse_historical_workbook")
    @mock.patch("apps.workflow.views.validate_historical_rows")
    def test_post_no_file_returns_error_partial(self, mock_validate, mock_parse):
        self.client.force_login(self.audit_user)
        url = reverse("workflow:historical-import")
        response = self.client.post(url, {})
        self.assertEqual(response.status_code, 200)
        self.assertTemplateUsed(response, "workflow/partials/historical_import_preview.html")
        self.assertContains(response, "Aucun fichier Excel sélectionné")


# ── Tests : build_historical_import_template ──────────────────────────────────


class BuildHistoricalTemplateTests(HistoricalImportMixin, TestCase):

    def test_template_has_data_sheet(self):
        wb = build_historical_import_template()
        self.assertIn(DATA_SHEET_NAME, wb.sheetnames)

    def test_template_has_correct_headers(self):
        wb = build_historical_import_template()
        ws = wb[DATA_SHEET_NAME]
        header_row = [cell.value for cell in ws[1]]
        for h in HEADERS:
            self.assertIn(h, header_row)
